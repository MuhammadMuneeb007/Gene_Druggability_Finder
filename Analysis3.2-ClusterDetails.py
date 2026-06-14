#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Analysis3.1-SupplementaryData2_DatasetInformation.py

Create Supplementary Data 2: dataset information workbook.

This script reads every generated dataset folder from:

    Datasets/
        dataset_catalogue.csv
        Dataset001_F01_T1/
            X_train.csv
            y_train.csv
            meta.json
            feature_columns.txt
            missingness_summary.csv

and produces:

    Supplementary_Data_2_Dataset_Information.xlsx

Main sheets:
    01_README
    02_Run_Summary
    03_Dataset_Catalogue
    04_Target_Distribution
    05_Feature_Block_Counts
    06_Feature_Group_Definitions
    07_Target_Definitions
    08_Missingness_Summary
    09_Feature_Columns_Long
    10_Audit_Checks

Run:
    python Analysis3.1-SupplementaryData2_DatasetInformation.py

Optional:
    python Analysis3.1-SupplementaryData2_DatasetInformation.py --datasets-dir Datasets
    python Analysis3.1-SupplementaryData2_DatasetInformation.py --out Supplementary_Data_2_Dataset_Information.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd


# =============================================================================
# DEFAULTS
# =============================================================================

DEFAULT_DATASETS_DIR = Path("Datasets")
DEFAULT_OUT_XLSX = Path("Supplementary_Data_2_Dataset_Information.xlsx")
DEFAULT_OUTDIR = Path("Supplementary_Data_2_Dataset_Information")


# =============================================================================
# TARGET DEFINITIONS
# =============================================================================

TARGETS: Dict[str, Dict[str, str]] = {
    "T1": {
        "target_name": "Clinical Target",
        "target_column": "clinical_target_label",
        "description": "Approved or clinically established drug target; strictest target definition.",
    },
    "T2": {
        "target_name": "Clinical Investigation Target",
        "target_column": "clinical_investigation_label",
        "description": "Gene with clinical-investigation evidence, including phase 1--3 clinical trial support.",
    },
    "T3": {
        "target_name": "Small-Molecule Target",
        "target_column": "small_molecule_druggable_label",
        "description": "Gene with small-molecule tractability evidence.",
    },
    "T4": {
        "target_name": "Chemical Tractability Target",
        "target_column": "chemical_tractable_label",
        "description": "Gene with broader chemical tractability evidence.",
    },
    "T5": {
        "target_name": "Biologic/Modality Target",
        "target_column": "biologic_druggable_label",
        "description": "Gene with biologic, antibody, targeted-degradation or other modality evidence.",
    },
    "T6": {
        "target_name": "Drug--Gene Interaction Target",
        "target_column": "dgidb_interaction_supported_label",
        "description": "Gene supported by DGIdb drug--gene interaction evidence.",
    },
    "T7": {
        "target_name": "Potentially Druggable Family Target",
        "target_column": "potentially_druggable_category_label",
        "description": "Gene assigned to a potentially druggable protein-family or category-based annotation.",
    },
    "T8": {
        "target_name": "Broad Druggability Target",
        "target_column": "final_any_druggable_label",
        "description": "Broad evidence-union label; positive if a gene satisfied any druggability definition.",
    },
}


# =============================================================================
# FEATURE BLOCK DEFINITIONS
# =============================================================================

FEATURE_BLOCKS: Dict[int, Dict[str, str]] = {
    1:  {"block": "F01", "name": "DepMap functional genomics", "domain": "Functional genomics / dependency"},
    2:  {"block": "F02", "name": "STRING protein--protein interaction network", "domain": "Network"},
    3:  {"block": "F03", "name": "Reactome biological pathways", "domain": "Pathway"},
    4:  {"block": "F04", "name": "Protein 3D structure", "domain": "Structure"},
    5:  {"block": "F05", "name": "InterPro/Pfam protein domains", "domain": "Protein domains"},
    6:  {"block": "F06", "name": "UniProt protein annotation", "domain": "Protein annotation"},
    7:  {"block": "F07", "name": "GTEx tissue expression", "domain": "Expression"},
    8:  {"block": "F08", "name": "Ensembl gene structure", "domain": "Gene structure"},
    9:  {"block": "F09", "name": "Constraint and disorder", "domain": "Constraint / disorder"},
    10: {"block": "F10", "name": "fpocket binding-pocket geometry", "domain": "Binding pockets"},
    11: {"block": "F11", "name": "Protein sequence composition", "domain": "Sequence / physicochemical"},
    12: {"block": "F12", "name": "ProtT5 protein embeddings", "domain": "Protein language model"},
    13: {"block": "F13", "name": "GWAS Catalog association burden", "domain": "Human genetics"},
    14: {"block": "F14", "name": "Gene Ontology annotation", "domain": "Functional annotation"},
    15: {"block": "F15", "name": "BioGRID curated interactions", "domain": "Curated network"},
    16: {"block": "F16", "name": "Human Protein Atlas", "domain": "Expression / localisation"},
    17: {"block": "F17", "name": "CTD chemical--gene interaction burden", "domain": "Chemical perturbation burden"},
    18: {"block": "F18", "name": "MGI mouse phenotype evidence", "domain": "Mouse phenotype"},
    19: {"block": "F19", "name": "gnomAD full constraint", "domain": "Population constraint"},
    20: {"block": "F20", "name": "CORUM protein complexes", "domain": "Protein complexes"},
    21: {"block": "F21", "name": "PhosphoSitePlus PTM sites", "domain": "Post-translational modification"},
    22: {"block": "F22", "name": "Ensembl paralogues", "domain": "Paralogues"},
}


# =============================================================================
# FEATURE GROUP DEFINITIONS
# =============================================================================

INDIVIDUAL_GROUPS = {f"F{n:02d}": [n] for n in range(1, 23)}
CUMULATIVE_GROUPS = {f"CUM{n:02d}": list(range(1, n + 1)) for n in range(1, 23)}
THEMATIC_GROUPS = {
    "GRP_Structure": [4, 10],
    "GRP_Network": [2, 15],
    "GRP_Expression": [7, 16],
    "GRP_Constraint": [9, 19],
    "GRP_Pathway": [3, 14],
    "GRP_Functional": [5, 6, 11, 12],
    "GRP_Omics": [1, 18],
    "GRP_Annotation": [8, 13, 20, 21, 22],
    "GRP_Literature": [17],
    "GRP_NoStructure": [n for n in range(1, 23) if n not in (4, 10)],
    "GRP_Full": list(range(1, 23)),
}

ALL_GROUPS: Dict[str, List[int]] = {}
ALL_GROUPS.update(INDIVIDUAL_GROUPS)
ALL_GROUPS.update(CUMULATIVE_GROUPS)
ALL_GROUPS.update(THEMATIC_GROUPS)


# =============================================================================
# HELPERS
# =============================================================================

def log(msg: str) -> None:
    print(msg, flush=True)


def now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def mkdir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def read_json(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text())
    except Exception:
        return {}


def safe_int(x: Any, default: int = 0) -> int:
    try:
        if x is None or pd.isna(x):
            return default
        return int(float(x))
    except Exception:
        return default


def safe_float(x: Any, default: float = np.nan) -> float:
    try:
        if x is None or pd.isna(x):
            return default
        return float(x)
    except Exception:
        return default


def extract_dataset_id(dataset_name: str) -> Optional[int]:
    m = re.match(r"Dataset(\d+)_", dataset_name)
    if not m:
        return None
    return int(m.group(1))


def extract_group_and_target(dataset_name: str) -> tuple[str, str]:
    """
    Examples:
        Dataset001_F01_T1              -> F01, T1
        Dataset177_CUM01_T1            -> CUM01, T1
        Dataset353_GRP_Structure_T1    -> GRP_Structure, T1
    """
    m = re.match(r"Dataset\d+_(.+)_(T\d+)$", dataset_name)
    if not m:
        return "", ""
    return m.group(1), m.group(2)


def feature_number_from_col(col: str) -> Optional[int]:
    m = re.match(r"^Feature(\d+)_", str(col))
    if not m:
        return None
    return int(m.group(1))


def classify_group(group_name: str) -> str:
    if group_name.startswith("F") and not group_name.startswith("Full"):
        return "Individual"
    if group_name.startswith("CUM"):
        return "Cumulative"
    if group_name.startswith("GRP"):
        return "Thematic"
    return "Unknown"


def list_dataset_dirs(datasets_dir: Path) -> List[Path]:
    dirs = [
        p for p in datasets_dir.iterdir()
        if p.is_dir() and re.match(r"^Dataset\d+_", p.name)
    ]
    return sorted(dirs, key=lambda p: extract_dataset_id(p.name) or 10**9)


def read_feature_columns(dataset_dir: Path) -> List[str]:
    txt = dataset_dir / "feature_columns.txt"
    xfile = dataset_dir / "X_train.csv"

    if txt.exists():
        return [x.strip() for x in txt.read_text().splitlines() if x.strip()]

    if xfile.exists():
        try:
            header = pd.read_csv(xfile, nrows=0).columns.tolist()
            return [c for c in header if c != "gene_symbol"]
        except Exception:
            return []

    return []


def read_target_distribution(dataset_dir: Path, target_col_from_meta: str = "") -> Dict[str, Any]:
    yfile = dataset_dir / "y_train.csv"

    out = {
        "target_column": target_col_from_meta,
        "n_genes_y": np.nan,
        "n_positive": np.nan,
        "n_negative": np.nan,
        "prevalence": np.nan,
    }

    if not yfile.exists():
        return out

    try:
        y = pd.read_csv(yfile, low_memory=False)
    except Exception:
        return out

    if "gene_symbol" in y.columns:
        possible_targets = [c for c in y.columns if c != "gene_symbol"]
    else:
        possible_targets = list(y.columns)

    if target_col_from_meta and target_col_from_meta in y.columns:
        target_col = target_col_from_meta
    elif possible_targets:
        target_col = possible_targets[0]
    else:
        return out

    yy = pd.to_numeric(y[target_col], errors="coerce")
    n = int(yy.notna().sum())
    pos = int((yy == 1).sum())
    neg = int((yy == 0).sum())
    prev = round(pos / n, 6) if n > 0 else np.nan

    out.update({
        "target_column": target_col,
        "n_genes_y": n,
        "n_positive": pos,
        "n_negative": neg,
        "prevalence": prev,
    })
    return out


def read_x_shape(dataset_dir: Path) -> Dict[str, Any]:
    xfile = dataset_dir / "X_train.csv"
    out = {
        "n_rows_x": np.nan,
        "n_columns_x": np.nan,
        "n_features_x": np.nan,
    }

    if not xfile.exists():
        return out

    try:
        # Count rows without loading full data.
        with open(xfile, "r", encoding="utf-8", errors="ignore") as f:
            n_lines = sum(1 for _ in f)
        header = pd.read_csv(xfile, nrows=0).columns.tolist()

        out["n_rows_x"] = max(n_lines - 1, 0)
        out["n_columns_x"] = len(header)
        out["n_features_x"] = len([c for c in header if c != "gene_symbol"])
    except Exception:
        pass

    return out


def read_missingness(dataset_dir: Path) -> Dict[str, Any]:
    mfile = dataset_dir / "missingness_summary.csv"

    out = {
        "mean_feature_missing_pct": np.nan,
        "median_feature_missing_pct": np.nan,
        "max_feature_missing_pct": np.nan,
        "n_features_missing_gt_50pct": np.nan,
        "n_features_missing_gt_90pct": np.nan,
    }

    if not mfile.exists():
        return out

    try:
        m = pd.read_csv(mfile)
    except Exception:
        return out

    if "pct_missing" not in m.columns:
        return out

    pct = pd.to_numeric(m["pct_missing"], errors="coerce")

    out.update({
        "mean_feature_missing_pct": round(float(pct.mean()), 6),
        "median_feature_missing_pct": round(float(pct.median()), 6),
        "max_feature_missing_pct": round(float(pct.max()), 6),
        "n_features_missing_gt_50pct": int((pct > 50).sum()),
        "n_features_missing_gt_90pct": int((pct > 90).sum()),
    })

    return out


def block_counts_from_features(feature_cols: List[str]) -> Dict[int, int]:
    counts = {i: 0 for i in range(1, 23)}
    for col in feature_cols:
        n = feature_number_from_col(col)
        if n is not None and 1 <= n <= 22:
            counts[n] += 1
    return counts


def blocks_present_from_counts(counts: Dict[int, int]) -> List[int]:
    return [i for i in range(1, 23) if counts.get(i, 0) > 0]


def make_csv_safe(df: pd.DataFrame) -> pd.DataFrame:
    """
    Convert lists/dicts to strings so CSV/XLSX export is clean.
    """
    out = df.copy()
    for col in out.columns:
        out[col] = out[col].map(
            lambda x: json.dumps(x) if isinstance(x, (list, dict)) else x
        )
    return out


# =============================================================================
# BUILD TABLES
# =============================================================================

def build_dataset_tables(datasets_dir: Path) -> Dict[str, pd.DataFrame]:
    dataset_dirs = list_dataset_dirs(datasets_dir)
    if not dataset_dirs:
        raise RuntimeError(f"No dataset folders found in {datasets_dir}")

    log(f"[DATASET DIRS FOUND] {len(dataset_dirs):,}")

    dataset_rows = []
    target_rows = []
    block_rows = []
    feature_long_rows = []
    missing_rows = []
    audit_rows = []

    for idx, ddir in enumerate(dataset_dirs, start=1):
        log(f"  [{idx:>4}/{len(dataset_dirs)}] {ddir.name}")

        meta = read_json(ddir / "meta.json")
        group_name, target_key = extract_group_and_target(ddir.name)

        dataset_id = meta.get("dataset_id", extract_dataset_id(ddir.name))
        dataset_name = meta.get("dataset_name", ddir.name)
        tier = meta.get("tier", classify_group(group_name))

        target_key = meta.get("target_key", target_key)
        target_column_meta = meta.get("target_column", "")
        target_name = meta.get("target_name", TARGETS.get(target_key, {}).get("target_name", ""))

        feature_cols = read_feature_columns(ddir)
        block_counts = block_counts_from_features(feature_cols)
        blocks_present = blocks_present_from_counts(block_counts)

        ydist = read_target_distribution(ddir, target_col_from_meta=target_column_meta)
        xshape = read_x_shape(ddir)
        miss = read_missingness(ddir)

        requested_blocks = meta.get("requested_feature_numbers", ALL_GROUPS.get(group_name, []))
        used_blocks = meta.get("feature_numbers_used", blocks_present)

        dataset_row = {
            "dataset_id": dataset_id,
            "dataset_name": dataset_name,
            "dataset_folder": str(ddir),
            "tier": tier,
            "group_name": group_name,
            "target_key": target_key,
            "target_name": target_name,
            "target_column": ydist.get("target_column", target_column_meta),
            "requested_feature_blocks": ",".join([f"F{x:02d}" for x in requested_blocks]) if isinstance(requested_blocks, list) else str(requested_blocks),
            "used_feature_blocks": ",".join([f"F{x:02d}" for x in used_blocks]) if isinstance(used_blocks, list) else str(used_blocks),
            "n_requested_feature_blocks": len(requested_blocks) if isinstance(requested_blocks, list) else np.nan,
            "n_used_feature_blocks": len(used_blocks) if isinstance(used_blocks, list) else np.nan,
            "n_features": xshape.get("n_features_x", meta.get("n_features_used", meta.get("n_features", np.nan))),
            "n_genes": ydist.get("n_genes_y", meta.get("n_genes", np.nan)),
            "n_positive": ydist.get("n_positive", meta.get("n_positives", np.nan)),
            "n_negative": ydist.get("n_negative", meta.get("n_negatives", np.nan)),
            "prevalence": ydist.get("prevalence", meta.get("prevalence", np.nan)),
            "mean_feature_missing_pct": miss["mean_feature_missing_pct"],
            "median_feature_missing_pct": miss["median_feature_missing_pct"],
            "max_feature_missing_pct": miss["max_feature_missing_pct"],
            "n_features_missing_gt_50pct": miss["n_features_missing_gt_50pct"],
            "n_features_missing_gt_90pct": miss["n_features_missing_gt_90pct"],
            "created_at": meta.get("created_at", ""),
        }
        dataset_rows.append(dataset_row)

        target_rows.append({
            "dataset_id": dataset_id,
            "dataset_name": dataset_name,
            "target_key": target_key,
            "target_name": target_name,
            "target_column": ydist.get("target_column", target_column_meta),
            "n_genes": ydist.get("n_genes_y", np.nan),
            "n_positive": ydist.get("n_positive", np.nan),
            "n_negative": ydist.get("n_negative", np.nan),
            "prevalence": ydist.get("prevalence", np.nan),
        })

        for block_num in range(1, 23):
            info = FEATURE_BLOCKS[block_num]
            n_features = block_counts.get(block_num, 0)
            block_rows.append({
                "dataset_id": dataset_id,
                "dataset_name": dataset_name,
                "tier": tier,
                "group_name": group_name,
                "target_key": target_key,
                "block": info["block"],
                "feature_block_number": block_num,
                "feature_block_name": info["name"],
                "biological_domain": info["domain"],
                "included": int(n_features > 0),
                "n_features_from_block": n_features,
            })

        for col in feature_cols:
            block_num = feature_number_from_col(col)
            info = FEATURE_BLOCKS.get(block_num, {})
            feature_long_rows.append({
                "dataset_id": dataset_id,
                "dataset_name": dataset_name,
                "tier": tier,
                "group_name": group_name,
                "target_key": target_key,
                "feature_column": col,
                "feature_block_number": block_num,
                "block": info.get("block", ""),
                "feature_block_name": info.get("name", ""),
                "biological_domain": info.get("domain", ""),
            })

        mfile = ddir / "missingness_summary.csv"
        if mfile.exists():
            try:
                mdf = pd.read_csv(mfile)
                if not mdf.empty:
                    for _, r in mdf.iterrows():
                        col = r.get("feature_column", "")
                        block_num = feature_number_from_col(col)
                        info = FEATURE_BLOCKS.get(block_num, {})
                        missing_rows.append({
                            "dataset_id": dataset_id,
                            "dataset_name": dataset_name,
                            "tier": tier,
                            "group_name": group_name,
                            "target_key": target_key,
                            "feature_column": col,
                            "feature_block_number": block_num,
                            "block": info.get("block", ""),
                            "n_missing": r.get("n_missing", np.nan),
                            "pct_missing": r.get("pct_missing", np.nan),
                            "n_nonmissing": r.get("n_nonmissing", np.nan),
                        })
            except Exception:
                pass

        audit_rows.append({
            "dataset_id": dataset_id,
            "dataset_name": dataset_name,
            "has_X_train": int((ddir / "X_train.csv").exists()),
            "has_y_train": int((ddir / "y_train.csv").exists()),
            "has_meta_json": int((ddir / "meta.json").exists()),
            "has_feature_columns": int((ddir / "feature_columns.txt").exists()),
            "has_missingness_summary": int((ddir / "missingness_summary.csv").exists()),
            "x_rows_equal_y_rows": int(
                pd.notna(xshape.get("n_rows_x")) and
                pd.notna(ydist.get("n_genes_y")) and
                int(xshape.get("n_rows_x")) == int(ydist.get("n_genes_y"))
            ),
            "n_rows_x": xshape.get("n_rows_x", np.nan),
            "n_rows_y": ydist.get("n_genes_y", np.nan),
            "n_features_x": xshape.get("n_features_x", np.nan),
            "n_features_listed": len(feature_cols),
        })

    tables = {
        "dataset_summary": pd.DataFrame(dataset_rows),
        "target_distribution": pd.DataFrame(target_rows),
        "feature_block_counts": pd.DataFrame(block_rows),
        "feature_columns_long": pd.DataFrame(feature_long_rows),
        "missingness_summary": pd.DataFrame(missing_rows),
        "audit_checks": pd.DataFrame(audit_rows),
    }

    return tables


def build_group_definitions_table() -> pd.DataFrame:
    rows = []
    for group_name, blocks in ALL_GROUPS.items():
        tier = classify_group(group_name)
        rows.append({
            "group_name": group_name,
            "tier": tier,
            "feature_blocks": ",".join([f"F{x:02d}" for x in blocks]),
            "n_feature_blocks": len(blocks),
            "feature_block_names": "; ".join([FEATURE_BLOCKS[x]["name"] for x in blocks if x in FEATURE_BLOCKS]),
        })
    return pd.DataFrame(rows)


def build_target_definitions_table() -> pd.DataFrame:
    rows = []
    for target_key, info in TARGETS.items():
        rows.append({
            "target_key": target_key,
            "target_name": info["target_name"],
            "target_column": info["target_column"],
            "description": info["description"],
        })
    return pd.DataFrame(rows)


def build_feature_block_definitions_table() -> pd.DataFrame:
    rows = []
    for n, info in FEATURE_BLOCKS.items():
        rows.append({
            "feature_block_number": n,
            "block": info["block"],
            "feature_block_name": info["name"],
            "biological_domain": info["domain"],
        })
    return pd.DataFrame(rows)


def build_run_summary_table(tables: Dict[str, pd.DataFrame]) -> pd.DataFrame:
    ds = tables["dataset_summary"]

    rows = [
        {"metric": "Created at", "value": now_iso()},
        {"metric": "Total datasets", "value": int(len(ds))},
        {"metric": "Unique feature groups", "value": int(ds["group_name"].nunique()) if "group_name" in ds else 0},
        {"metric": "Unique targets", "value": int(ds["target_key"].nunique()) if "target_key" in ds else 0},
        {"metric": "Mean genes per dataset", "value": round(float(pd.to_numeric(ds["n_genes"], errors="coerce").mean()), 3)},
        {"metric": "Mean features per dataset", "value": round(float(pd.to_numeric(ds["n_features"], errors="coerce").mean()), 3)},
        {"metric": "Minimum features in a dataset", "value": int(pd.to_numeric(ds["n_features"], errors="coerce").min())},
        {"metric": "Maximum features in a dataset", "value": int(pd.to_numeric(ds["n_features"], errors="coerce").max())},
        {"metric": "Mean positives per dataset", "value": round(float(pd.to_numeric(ds["n_positive"], errors="coerce").mean()), 3)},
        {"metric": "Mean target prevalence", "value": round(float(pd.to_numeric(ds["prevalence"], errors="coerce").mean()), 6)},
    ]
    return pd.DataFrame(rows)


def build_readme_table() -> pd.DataFrame:
    rows = [
        {
            "sheet": "01_README",
            "description": "Description of Supplementary Data 2 workbook.",
        },
        {
            "sheet": "02_Run_Summary",
            "description": "Overall summary of generated datasets, targets, feature counts and prevalence.",
        },
        {
            "sheet": "03_Dataset_Catalogue",
            "description": "One row per generated dataset, including dataset ID, feature group, target, number of genes, positives, negatives, prevalence and feature counts.",
        },
        {
            "sheet": "04_Target_Distribution",
            "description": "Target distribution for each dataset.",
        },
        {
            "sheet": "05_Feature_Block_Counts",
            "description": "Feature block inclusion and number of features contributed by each block for each dataset.",
        },
        {
            "sheet": "06_Feature_Group_Definitions",
            "description": "Definitions of individual, cumulative and thematic feature groups.",
        },
        {
            "sheet": "07_Target_Definitions",
            "description": "Definitions of the eight druggability target labels.",
        },
        {
            "sheet": "08_Missingness_Summary",
            "description": "Feature-level missingness summaries when missingness_summary.csv files are available.",
        },
        {
            "sheet": "09_Feature_Columns_Long",
            "description": "Long-format list of all feature columns used by each dataset.",
        },
        {
            "sheet": "10_Audit_Checks",
            "description": "Checks for required files and consistency between X and y row counts.",
        },
    ]
    return pd.DataFrame(rows)


# =============================================================================
# SAVE OUTPUTS
# =============================================================================

def save_csv_outputs(tables: Dict[str, pd.DataFrame], outdir: Path) -> None:
    mkdir(outdir)
    for name, df in tables.items():
        path = outdir / f"{name}.csv"
        make_csv_safe(df).to_csv(path, index=False)


def autosize_worksheet(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame, max_width: int = 60) -> None:
    """
    Works with openpyxl engine.
    """
    ws = writer.sheets[sheet_name]
    for idx, col in enumerate(df.columns, start=1):
        series = df[col].astype(str).replace("nan", "")
        max_len = max([len(str(col))] + [len(x) for x in series.head(1000).tolist()])
        width = min(max(max_len + 2, 10), max_width)
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def save_excel(tables: Dict[str, pd.DataFrame], out_xlsx: Path) -> None:
    """
    Save polished workbook using pandas + openpyxl.
    """
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        sheet_map = [
            ("01_README", tables["readme"]),
            ("02_Run_Summary", tables["run_summary"]),
            ("03_Dataset_Catalogue", tables["dataset_summary"]),
            ("04_Target_Distribution", tables["target_distribution"]),
            ("05_Feature_Block_Counts", tables["feature_block_counts"]),
            ("06_Feature_Group_Definitions", tables["feature_group_definitions"]),
            ("07_Target_Definitions", tables["target_definitions"]),
            ("08_Missingness_Summary", tables["missingness_summary"]),
            ("09_Feature_Columns_Long", tables["feature_columns_long"]),
            ("10_Audit_Checks", tables["audit_checks"]),
            ("11_Feature_Block_Definitions", tables["feature_block_definitions"]),
        ]

        for sheet_name, df in sheet_map:
            df2 = make_csv_safe(df)
            # Excel sheet limit is 1,048,576 rows. Protect against huge long table.
            if len(df2) > 1_048_000:
                df2 = df2.head(1_048_000).copy()
                df2.loc[len(df2)] = ["TRUNCATED_DUE_TO_EXCEL_ROW_LIMIT"] + [""] * (df2.shape[1] - 1)
            df2.to_excel(writer, sheet_name=sheet_name, index=False)

        # Formatting
        wb = writer.book
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        header_font = Font(bold=True)
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for sheet_name, df in sheet_map:
            ws = writer.sheets[sheet_name]

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = border

            autosize_worksheet(writer, sheet_name, df)

        # Number formats for prevalence / percentages
        for sheet_name in ["03_Dataset_Catalogue", "04_Target_Distribution"]:
            ws = writer.sheets[sheet_name]
            header = [c.value for c in ws[1]]
            for col_name in ["prevalence"]:
                if col_name in header:
                    col_idx = header.index(col_name) + 1
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row=row, column=col_idx).number_format = "0.0000"

        for sheet_name in ["03_Dataset_Catalogue", "08_Missingness_Summary"]:
            ws = writer.sheets[sheet_name]
            header = [c.value for c in ws[1]]
            for col_name in ["mean_feature_missing_pct", "median_feature_missing_pct", "max_feature_missing_pct", "pct_missing"]:
                if col_name in header:
                    col_idx = header.index(col_name) + 1
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row=row, column=col_idx).number_format = "0.00"


# =============================================================================
# MAIN
# =============================================================================

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create Supplementary Data 2 dataset information workbook."
    )
    parser.add_argument(
        "--datasets-dir",
        default=str(DEFAULT_DATASETS_DIR),
        help=f"Directory containing generated dataset folders. Default: {DEFAULT_DATASETS_DIR}",
    )
    parser.add_argument(
        "--out",
        default=str(DEFAULT_OUT_XLSX),
        help=f"Output Excel workbook. Default: {DEFAULT_OUT_XLSX}",
    )
    parser.add_argument(
        "--csv-outdir",
        default=str(DEFAULT_OUTDIR),
        help=f"Directory for CSV copies of workbook sheets. Default: {DEFAULT_OUTDIR}",
    )

    args = parser.parse_args()

    datasets_dir = Path(args.datasets_dir)
    out_xlsx = Path(args.out)
    outdir = Path(args.csv_outdir)

    if not datasets_dir.exists():
        raise FileNotFoundError(f"Datasets directory not found: {datasets_dir}")

    log("=" * 100)
    log("SUPPLEMENTARY DATA 2 — DATASET INFORMATION")
    log(f"Started:       {now_iso()}")
    log(f"Datasets dir:  {datasets_dir.resolve()}")
    log(f"Output XLSX:   {out_xlsx.resolve()}")
    log(f"CSV outdir:    {outdir.resolve()}")
    log("=" * 100)

    tables = build_dataset_tables(datasets_dir)
    tables["readme"] = build_readme_table()
    tables["feature_group_definitions"] = build_group_definitions_table()
    tables["target_definitions"] = build_target_definitions_table()
    tables["feature_block_definitions"] = build_feature_block_definitions_table()
    tables["run_summary"] = build_run_summary_table(tables)

    save_csv_outputs(tables, outdir)
    save_excel(tables, out_xlsx)

    log("\n" + "=" * 100)
    log("[DONE]")
    log(f"Saved Excel: {out_xlsx.resolve()}")
    log(f"Saved CSV sheets: {outdir.resolve()}")
    log("=" * 100)

    # Print compact summary.
    ds = tables["dataset_summary"]
    log(f"Total datasets: {len(ds):,}")
    log(f"Unique groups:  {ds['group_name'].nunique():,}")
    log(f"Unique targets: {ds['target_key'].nunique():,}")
    log(f"Mean features:  {pd.to_numeric(ds['n_features'], errors='coerce').mean():.2f}")
    log(f"Mean genes:     {pd.to_numeric(ds['n_genes'], errors='coerce').mean():.2f}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        log("\n[STOPPED] Interrupted by user.")
        sys.exit(130)
    except Exception as exc:
        log(f"\n[ERROR] {exc}")
        raise